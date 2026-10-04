"""Converts docs/IEEE_REPORT.md (a small, regular Markdown subset) to a Word file.

    ml/.venv/bin/python docs/md2docx.py docs/IEEE_REPORT.md docs/IEEE_REPORT.docx

Needs python-docx (`pip install python-docx`).
"""
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

src, out = sys.argv[1], sys.argv[2]
lines = open(src, encoding="utf-8").read().split("\n")

doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Cm(2.2)
sec.top_margin = sec.bottom_margin = Cm(2.0)
base = doc.styles["Normal"]
base.font.name = "Times New Roman"
base.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
base.font.size = Pt(11)
base.paragraph_format.space_after = Pt(6)
for name, size in (("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)):
    st = doc.styles[name]
    st.font.name = "Times New Roman"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor(0x1F, 0x2A, 0x44)

TOKEN = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*\s][^*]*?\*)")


def add_runs(par, text, size=None):
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = par.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt((size or 11) - 1)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r = par.add_run(part[1:-1])
            r.italic = True
        else:
            r = par.add_run(part)
        if size and not part.startswith("`"):
            r.font.size = Pt(size)


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def add_table(rows):
    header, body = rows[0], rows[1:]
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        add_runs(c.paragraphs[0], h.strip(), size=9)
        for r in c.paragraphs[0].runs:
            r.bold = True
        shade(c, "E4E9F2")
    for row in body:
        cells = t.add_row().cells
        for i in range(len(header)):
            cells[i].text = ""
            add_runs(cells[i].paragraphs[0], (row[i] if i < len(row) else "").strip(), size=9)
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(1)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


i = 0
while i < len(lines):
    line = lines[i]
    if not line.strip() or line.strip() == "---":
        i += 1
        continue
    if line.startswith("|"):
        block = []
        while i < len(lines) and lines[i].startswith("|"):
            cells = re.split(r"(?<!\\)\|", lines[i].strip())[1:-1]
            if not all(re.fullmatch(r"\s*:?-{3,}:?\s*", c) for c in cells):
                block.append(cells)
            i += 1
        add_table(block)
        continue
    m = re.match(r"^(#{1,3})\s+(.*)", line)
    if m:
        doc.add_heading(m.group(2), level=len(m.group(1)))
        i += 1
        continue
    m = re.match(r"^\s*- (.*)", line)
    if m:
        add_runs(doc.add_paragraph(style="List Bullet"), m.group(1))
        i += 1
        continue
    m = re.match(r"^\d+\.\s+(.*)", line)
    if m:
        add_runs(doc.add_paragraph(style="List Number"), m.group(1))
        i += 1
        continue
    m = re.match(r"^\[(\d+)\]\s+(.*)", line)
    if m:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(1.0)
        p.paragraph_format.first_line_indent = Cm(-1.0)
        add_runs(p, f"[{m.group(1)}]  {m.group(2)}", size=10)
        i += 1
        continue
    buf = [line.strip()]
    i += 1
    while i < len(lines) and lines[i].strip() and not re.match(r"^(#{1,3}\s|\||\s*- |\d+\.\s|\[\d+\]|---)", lines[i]):
        buf.append(lines[i].strip())
        i += 1
    add_runs(doc.add_paragraph(), " ".join(buf))

doc.save(out)
print("saved", out)
